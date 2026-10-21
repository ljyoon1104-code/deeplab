from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement as E
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[3]
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.55); sec.bottom_margin=Cm(1.45)
sec.left_margin=Cm(1.7); sec.right_margin=Cm(1.7)
sec.footer_distance=Cm(.65)
for name in ['Normal','Title','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='맑은 고딕'; st.font.color.rgb=RGBColor(0,0,0)
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'맑은 고딕')
    st.paragraph_format.space_after=Pt(6)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.line_spacing=1.1
doc.styles['Title'].font.size=Pt(21)
doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(0)
doc.styles['Heading 2'].font.size=Pt(12)
doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)

class Math:
    def __init__(self, nodes): self.nodes=nodes

def R(s, upright=False):
    r=E('m:r')
    if upright:
        pr=E('m:rPr'); sty=E('m:sty'); sty.set(qn('m:val'),'p'); pr.append(sty); r.append(pr)
    rp=E('w:rPr'); fonts=E('w:rFonts')
    for a in ['ascii','hAnsi']: fonts.set(qn('w:'+a),'Cambria Math')
    rp.append(fonts); sz=E('w:sz'); sz.set(qn('w:val'),'22'); rp.append(sz); r.append(rp)
    t=E('m:t'); t.text=str(s); t.set(qn('xml:space'),'preserve'); r.append(t)
    return Math([r])

def M(*parts):
    nodes=[]
    for p in parts:
        nodes.extend(deepcopy(p.nodes if isinstance(p,Math) else R(p).nodes))
    return Math(nodes)

def holder(tag, val):
    e=E('m:'+tag)
    for n in M(val).nodes: e.append(n)
    return e

def F(a,b):
    e=E('m:f'); e.append(holder('num',a)); e.append(holder('den',b)); return Math([e])

def S(a,b):
    e=E('m:sSub'); e.append(holder('e',a)); e.append(holder('sub',b)); return Math([e])

def U(a,b):
    e=E('m:sSup'); e.append(holder('e',a)); e.append(holder('sup',b)); return Math([e])

def D(a,b): return F(M('∂',a),M('∂',b))
def ln(a): return M(R('ln',True),'(',a,')')
def sig(a): return M('σ(',a,')')
def relu(a): return M(R('ReLU',True),'(',a,')')

def Sum(sub,sup,body):
    n=E('m:nary'); pr=E('m:naryPr')
    ch=E('m:chr'); ch.set(qn('m:val'),'∑'); pr.append(ch)
    loc=E('m:limLoc'); loc.set(qn('m:val'),'undOvr'); pr.append(loc); n.append(pr)
    n.append(holder('sub',sub)); n.append(holder('sup',sup)); n.append(holder('e',body)); return Math([n])

def put(p,m):
    eq=E('m:oMath')
    for n in M(m).nodes: eq.append(n)
    p._p.append(eq)

def P(*parts,bold=False,size=None):
    p=doc.add_paragraph()
    for x in parts:
        if isinstance(x,Math): put(p,x)
        else:
            r=p.add_run(str(x)); r.bold=bold
            if size: r.font.size=Pt(size)
    return p

def EQ(*parts):
    p=P(M(*parts)); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(8)
    return p

def H(t): doc.add_heading(t,2)
def PAGE(t): doc.add_page_break(); doc.add_heading(t,1)

def T(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.autofit=False
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    pr=t._tbl.tblPr; borders=E('w:tblBorders')
    for k in ['top','bottom','left','right','insideH','insideV']:
        b=E('w:'+k); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'5'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    pr.append(borders)
    for i,values in enumerate([headers]+rows):
        row=t.rows[0] if i==0 else t.add_row()
        rp=row._tr.get_or_add_trPr(); rp.append(E('w:cantSplit'))
        if i==0: rp.append(E('w:tblHeader'))
        for j,v in enumerate(values):
            c=row.cells[j]; c.width=Cm(widths[j]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr(); margins=E('w:tcMar')
            for k,n in [('top',90),('bottom',90),('left',95),('right',95)]:
                e=E('w:'+k); e.set(qn('w:w'),str(n)); e.set(qn('w:type'),'dxa'); margins.append(e)
            cp.append(margins)
            shade=E('w:shd'); shade.set(qn('w:fill'),'DBE8F4' if i==0 else ('F7F9FB' if i%2==0 else 'FFFFFF')); cp.append(shade)
            p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
            if isinstance(v,Math):
                put(p,v)
            else:
                r=p.add_run(str(v)); r.font.size=Pt(10); r.bold=i==0
    p=P(); p.paragraph_format.line_spacing=Pt(1); p.paragraph_format.space_after=Pt(2); p.add_run().font.size=Pt(1)

fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=E('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); fp._p.append(field)

# The original chat examples, with the original values and explanations.
doc.add_paragraph('역전파 숫자 예제',style='Title')
P('자주 쓰는 조합별로, 미분 결과에 숫자만 넣으면 되는 표로 정리할게요.')
P('여기서 ',D('L','w'),'는 가중치 ',R('w'),'의 기울기예요.')
H('① 시그모이드와 BCE 정답이 1일 때')
P('입력 ',R('x = 2'),', 예측 확률 ',R('p = 0.5'),', 정답 ',R('y = 1'),'이라고 할게요.')
T(['계산 단계','가져다 쓰는 미분 공식','숫자 대입'],[
    [M('L = −',ln('p')),M(D('L','p'),' = −',F(1,'p')),M('−',F(1,'0.5'),' = −2')],
    [M('p = ',sig('z')),M(D('p','z'),' = p(1 − p)'),M('0.5 × (1 − 0.5) = 0.25')],
    [R('z = wx + b'),M(D('z','w'),' = x'),R('2')],
    ['최종 가중치 기울기','위 세 값을 곱하기',R('(−2) × 0.25 × 2 = −1')],
],[4.4,5.2,8.0])
P('기울기가 음수라서 업데이트하면 가중치가 커져요.')
H('② 시그모이드와 BCE 정답이 0일 때')
P('똑같이 ',R('x = 2'),', ',R('p = 0.5'),'이고, 정답만 ',R('y = 0'),'으로 바꿔볼게요.')
T(['계산 단계','가져다 쓰는 미분 공식','숫자 대입'],[
    [M('L = −',ln('1 − p')),M(D('L','p'),' = ',F(1,'1 − p')),M(F(1,'1 − 0.5'),' = 2')],
    [M('p = ',sig('z')),M(D('p','z'),' = p(1 − p)'),R('0.5 × (1 − 0.5) = 0.25')],
    [R('z = wx + b'),M(D('z','w'),' = x'),R('2')],
    ['최종 가중치 기울기','위 세 값을 곱하기',R('2 × 0.25 × 2 = 1')],
],[4.4,5.2,8.0])
P('이번에는 기울기가 양수라서 업데이트하면 가중치가 작아져요.')
P('①과 ②는 다음 공식 하나로 바로 계산할 수도 있어요.')
EQ(D('L','w'),' = (p − y)x')
P('이 공식은 시그모이드와 BCE를 함께 쓸 때 적용돼요.',bold=True)

PAGE('③ 소프트맥스와 CCE 여러 클래스의 경우')
P('입력은 ',R('x = 2'),', 정답은 고양이라고 해볼게요.')
EQ('y = [1, 0, 0],     p = [0.7, 0.2, 0.1]')
P('각 클래스에 점수 ',M(S('z','k'),' = ',S('w','k'),'x + ',S('b','k')),'를 만들고, 소프트맥스로 확률을 계산해요.')
P('소프트맥스와 CCE를 합쳐 미분한 결과는 다음과 같아요.')
EQ(D('L',S('z','k')),' = ',S('p','k'),' − ',S('y','k'))
T(['클래스',M('예측 ',S('p','k')),M('정답 ',S('y','k')),'점수의 기울기','가중치 기울기'],[
    ['고양이',R('0.7'),R('1'),R('0.7 − 1 = −0.3'),R('−0.3 × 2 = −0.6')],
    ['강아지',R('0.2'),R('0'),R('0.2 − 0 = 0.2'),R('0.2 × 2 = 0.4')],
    ['새',R('0.1'),R('0'),R('0.1 − 0 = 0.1'),R('0.1 × 2 = 0.2')],
],[2.1,2.4,2.1,5.1,5.9])
P('업데이트하면 정답인 고양이의 점수는 올라가고, 강아지와 새의 점수는 내려가는 방향으로 움직여요.')
P('위 표는 모두 데이터 하나 기준이에요. 여러 데이터의 평균 손실을 사용한다면, 각 데이터에서 구한 가중치 기울기도 평균 내서 업데이트에 사용해요.')
H('CCE 손실 계산도 함께 확인')
EQ('L = −',Sum('k = 1','C',M(S('y','k'),ln(S('p','k')))))
EQ('L = −[1 × ',ln('0.7'),' + 0 × ',ln('0.2'),' + 0 × ',ln('0.1'),']')
EQ('L = −',ln('0.7'),' ≈ 0.357')
P('정답 클래스에 준 확률의 ',M('−',ln('p')),'만 남아요.')

PAGE('④ 은닉층이 있을 때 순전파')
P('입력 2개 → ReLU 은닉 뉴런 2개 → 시그모이드 출력 1개 → BCE 손실인 예제를 볼게요.')
P('입력은 ',M(S('x',1),' = 1,  ',S('x',2),' = 2'),', 정답은 ',R('y = 1'),', 학습률은 ',R('η = 0.1'),'이에요.')
P('은닉층으로 들어가는 가중치는 ',R('w'),', 은닉층에서 출력층으로 연결되는 가중치는 ',R('v'),'라고 부를게요.')
H('1단계 처음 가중치와 편향을 정해요')
T(['위치','가중치와 편향의 초기값'],[
    ['은닉 뉴런 1',M(S('w','11'),' = 1,  ',S('w','12'),' = 0,  ',S('b',1),' = 0')],
    ['은닉 뉴런 2',M(S('w','21'),' = −1,  ',S('w','22'),' = 0,  ',S('b',2),' = 0')],
    ['출력 뉴런',M(S('v',1),' = 1,  ',S('v',2),' = 1,  ',S('b','o'),' = −1')],
],[4.0,13.6])
H('2단계 은닉층부터 순전파해요')
P(R('a'),'는 ReLU에 넣기 전 값, ',R('h'),'는 ReLU를 통과한 은닉 뉴런의 출력이에요.')
T(['구할 값','숫자를 넣은 계산','결과'],[
    [M('은닉 입력 ',S('a',1)),R('1 × 1 + 0 × 2 + 0'),R('1')],
    [M('은닉 출력 ',S('h',1)),relu(1),R('1')],
    [M('은닉 입력 ',S('a',2)),R('(−1) × 1 + 0 × 2 + 0'),R('−1')],
    [M('은닉 출력 ',S('h',2)),relu('−1'),R('0')],
    [R('출력층 점수 z'),R('1 × 1 + 1 × 0 − 1'),R('0')],
    [R('예측 확률 p'),M(sig(0),' = ',F(1,M('1 + ',U('e','0')))),R('0.5')],
    [R('BCE 손실 L'),M('−',ln('0.5')),R('0.6931')],
],[4.0,10.5,3.1])
P('ReLU는 양수는 그대로 보내고 음수는 0으로 바꿔요. 그래서 첫 번째 은닉 뉴런의 출력은 1, 두 번째는 0이에요.')

PAGE('⑤ 은닉층으로 기울기를 전달해요')
H('3단계 출력층에서 역전파를 시작해요')
P('시그모이드와 BCE 조합이므로 출력층 점수의 기울기는 ',R('p − y'),'예요.')
EQ(D('L','z'),' = p − y = 0.5 − 1 = −0.5')
T(['구할 기울기','미분 결과에 숫자 대입','결과'],[
    [M(D('L',S('v',1))),M(D('L','z'),' × ',S('h',1),' = −0.5 × 1'),R('−0.5')],
    [M(D('L',S('v',2))),M(D('L','z'),' × ',S('h',2),' = −0.5 × 0'),R('0')],
    [M(D('L',S('b','o'))),D('L','z'),R('−0.5')],
],[3.5,11.0,3.1])
H('4단계 출력층에서 은닉층으로 돌아가요')
P('출력층에서 돌아온 기울기에 연결 가중치 ',R('v'),'를 곱하고, ReLU의 기울기를 곱해요.')
EQ(D('L',S('a','j')),' = ',D('L','z'),' × ',S('v','j'),' × ',D(S('h','j'),S('a','j')))
T(['은닉 뉴런','ReLU의 기울기','은닉층 입력의 기울기'],[
    ['1',M(S('a',1),' = 1 > 0 → 1'),M(D('L',S('a',1)),' = −0.5 × 1 × 1 = −0.5')],
    ['2',M(S('a',2),' = −1 < 0 → 0'),M(D('L',S('a',2)),' = −0.5 × 1 × 0 = 0')],
],[2.2,5.1,10.3])
P('두 번째 은닉 뉴런은 ReLU 입력이 음수여서, 이 데이터에서는 뒤로 전달되는 기울기가 0이에요.')
H('5단계 입력값을 곱해 은닉층 가중치 기울기를 구해요')
T(['구할 기울기','미분 결과에 숫자 대입','결과'],[
    [D('L',S('w','11')),M(D('L',S('a',1)),' × ',S('x',1),' = −0.5 × 1'),R('−0.5')],
    [D('L',S('w','12')),M(D('L',S('a',1)),' × ',S('x',2),' = −0.5 × 2'),R('−1.0')],
    [D('L',S('b',1)),D('L',S('a',1)),R('−0.5')],
    [M(S('w','21'),', ',S('w','22'),', ',S('b',2)),'두 번째 은닉 뉴런의 기울기가 0이므로 모두 0',R('0')],
],[3.5,11.0,3.1])
P('은닉 뉴런이 여러 출력 뉴런에 연결되면, 각 연결에서 돌아오는 기울기를 더한 뒤 ReLU 기울기를 곱해요.',size=9.5)

PAGE('⑥ 가중치를 업데이트해요')
H('6단계 경사하강법으로 값을 바꿔요')
EQ(S('w','new'),' = w − η',D('L','w'))
P('학습률 ',R('η = 0.1'),'을 적용하면 다음과 같아요. 편향도 같은 방식으로 업데이트해요.')
T(['대상','현재 값','기울기','계산','새 값'],[
    [S('w','11'),R('1'),R('−0.5'),R('1 − 0.1 × (−0.5)'),R('1.05')],
    [S('w','12'),R('0'),R('−1'),R('0 − 0.1 × (−1)'),R('0.10')],
    [S('b',1),R('0'),R('−0.5'),R('0 − 0.1 × (−0.5)'),R('0.05')],
    [S('w','21'),R('−1'),R('0'),R('−1 − 0.1 × 0'),R('−1')],
    [S('w','22'),R('0'),R('0'),R('0 − 0.1 × 0'),R('0')],
    [S('b',2),R('0'),R('0'),R('0 − 0.1 × 0'),R('0')],
    [S('v',1),R('1'),R('−0.5'),R('1 − 0.1 × (−0.5)'),R('1.05')],
    [S('v',2),R('1'),R('0'),R('1 − 0.1 × 0'),R('1')],
    [S('b','o'),R('−1'),R('−0.5'),R('−1 − 0.1 × (−0.5)'),R('−0.95')],
],[2.0,2.2,2.2,8.0,3.2])
P('기울기는 모두 업데이트 전 가중치로 계산해요. 모든 기울기를 구한 다음 가중치와 편향을 바꿔요.',bold=True)
H('7단계 다시 순전파해서 확인해요')
EQ(S('h',1),' = ',relu('1.05 × 1 + 0.10 × 2 + 0.05'),' = 1.30')
EQ(S('h',2),' = ',relu('−1'),' = 0')
EQ('z = 1.05 × 1.30 + 1 × 0 − 0.95 = 0.415')
EQ('p = ',sig('0.415'),' ≈ 0.6023')
EQ('L = −',ln('0.6023'),' ≈ 0.5070')
P('정답이 1인데 예측 확률은 0.5 → 0.6023으로 올라갔고, 손실은 0.6931 → 0.5070으로 줄었어요.')
P('출력층뿐 아니라 은닉층의 가중치도 함께 바뀐 결과예요.')

for root in [doc.styles.element,doc.element]:
    for b in root.xpath('.//w:pBdr'): b.getparent().remove(b)
doc.core_properties.title='역전파 숫자 예제'
doc.core_properties.author=''
out=ROOT/'output'/'document-build'/'original_examples_typeset.docx'
out.parent.mkdir(parents=True, exist_ok=True)
doc.save(out)
print(out)
