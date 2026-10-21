from pathlib import Path
import math
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
OUT = ROOT / 'output' / '앱_역전파_숫자예제.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.4), Cm(1.4)
sec.left_margin, sec.right_margin = Cm(1.75), Cm(1.75)
sec.footer_distance = Cm(.7)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = '맑은 고딕'
    st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '맑은 고딕')
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(4)
doc.styles['Normal'].font.size = Pt(10)
doc.styles['Normal'].paragraph_format.line_spacing = 1.06
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(17)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 2'].font.size = Pt(11.5)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)

def para(text='', bold=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    return p

def heading(text):
    return doc.add_heading(text, level=2)

def equation(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    o = OxmlElement('m:oMath')
    r = OxmlElement('m:r')
    t = OxmlElement('m:t')
    t.text = text
    r.append(t)
    o.append(r)
    p._p.append(o)
    return p

def table(headers, rows, widths, font=9.3):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    props = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '5')
        el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    props.append(borders)
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
    for idx, values in enumerate([headers] + rows):
        cells = t.rows[0].cells if idx == 0 else t.add_row().cells
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if idx == 0:
            rep = OxmlElement('w:tblHeader'); trpr.append(rep)
        for col, value in enumerate(values):
            c = cells[col]
            c.width = Cm(widths[col])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = c._tc.get_or_add_tcPr()
            mar = OxmlElement('w:tcMar')
            for side, val in [('top',65),('bottom',65),('left',100),('right',100)]:
                m = OxmlElement('w:'+side); m.set(qn('w:w'), str(val)); m.set(qn('w:type'), 'dxa'); mar.append(m)
            cp.append(mar)
            sh = OxmlElement('w:shd')
            sh.set(qn('w:fill'), 'DBE8F4' if idx == 0 else ('F7F9FB' if idx % 2 == 0 else 'FFFFFF'))
            cp.append(sh)
            p = c.paragraphs[0]
            if re.fullmatch(r'[−\-]?\d+(?:\.\d+)?', str(value)):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            r = p.add_run(str(value))
            r.font.size = Pt(font)
            r.bold = idx == 0
            r.font.name = '맑은 고딕'
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '맑은 고딕')
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(1)
    spacer.add_run().font.size = Pt(1)
    return t

def newpage(title):
    doc.add_page_break()
    doc.add_heading(title, level=1)

f = sec.footer.paragraphs[0]
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run(); r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); f._p.append(fld)

# Page 1
doc.add_paragraph('앱에서 배우는 역전파\n숫자 예제와 공식표', style='Title')
para('미분 결과는 표에서 가져오고, 현재 숫자를 대입한 뒤 곱하면 됩니다. 순전파부터 손실 계산, 역전파, 가중치 갱신까지 순서대로 확인하세요.')
para('앱의 BCEE는 BCE, CCEE는 CCE와 같은 뜻입니다. 이 문서의 ln은 자연로그이며, 계산 결과는 소수점 넷째 자리까지 표시합니다.', size=9)
heading('1 앱에서 사용하는 계산')
table(['앱의 학습 내용', '사용하는 방식'], [
    ['3차시와 4차시', 'ReLU 은닉층과 시그모이드 또는 소프트맥스 출력의 역할'],
    ['5차시 손실 계산', 'MSE와 BCEE 및 CCEE의 숫자 계산'],
    ['6차시 가중치 갱신', '시그모이드 + BCEE의 기울기로 경사하강법 수행'],
    ['8차시 숫자 분류', 'ReLU 은닉층 + 소프트맥스 + CCEE를 사용하고 Adam으로 갱신'],
], [4.2,13.3])
heading('2 미분 없이 가져다 쓰는 공식')
para('g는 해당 값에 대한 손실의 기울기입니다. 예를 들어 g_w는 ∂L/∂w를 뜻해요.', size=9.5)
table(['계산', '가져다 쓰는 기울기 공식'], [
    ['가중합 z = Σ(wᵢxᵢ) + b', 'g_wᵢ = g_z × xᵢ     /     g_b = g_z'],
    ['시그모이드 p = σ(z)', '출력의 기울기 dp/dz = p(1 − p)'],
    ['시그모이드 + BCEE', 'g_z = p − y'],
    ['소프트맥스 + CCEE', '클래스 k마다 g_zₖ = pₖ − yₖ'],
    ['ReLU h = max(0, a)', 'a > 0이면 g_a = g_h, a < 0이면 g_a = 0'],
    ['경사하강법', '새 가중치 = 현재 가중치 − η × 가중치 기울기'],
], [5.5,12.0])
para('η는 학습률, y는 정답, p는 예측 확률입니다. ReLU는 a = 0에서 미분할 수 없으며 일반적인 학습 구현에서는 기울기를 0으로 처리합니다. 계단 함수는 보통의 역전파 학습용으로 사용하지 않습니다.', size=9)
heading('3 MSE는 손실 계산만 확인')
table(['앱의 데이터', '정답', '예측', '예측 − 정답', '차이의 제곱'], [
    ['학생 1','80','85','5','25'], ['학생 2','70','80','10','100'], ['학생 3','90','90','0','0'],
], [3.3,2,2,5.1,5.1], font=9)
para('MSE = (25 + 100 + 0) ÷ 3 = 41.6667입니다. 아래 분류 예제에서는 BCEE나 CCEE로 손실을 계산합니다.', size=9.2)

# Page 2
p0=1/(1+math.exp(-.1)); d=p0-1
w1=.2-.5*d; w2=-.1-.5*d; b=-.5*d
znew=w1+w2+b; pnew=1/(1+math.exp(-znew))
newpage('시그모이드와 BCEE 정답이 1인 경우')
para('6차시의 실제 초기값을 사용합니다. 입력 x₁ = 1, x₂ = 1, 정답 y = 1, 가중치 w₁ = 0.2, w₂ = −0.1, 편향 b = 0, 학습률 η = 0.5입니다.')
heading('1 순전파와 손실 계산')
table(['구할 값', '숫자를 넣은 계산', '결과'], [
    ['가중합 z','0.2 × 1 + (−0.1) × 1 + 0','0.1000'],
    ['예측 확률 p','1 ÷ (1 + exp(−0.1))',f'{p0:.4f}'],
    ['BCEE 손실 L','−ln(p)  정답이 1이므로 적용',f'{-math.log(p0):.4f}'],
], [3.2,10.8,3.5])
heading('2 역전파에서 단계별 기울기 계산')
table(['단계', '미분 결과에 숫자 대입', '결과'], [
    ['손실 → 확률','g_p = −1 ÷ p = −1 ÷ 0.524979…', f'{-1/p0:.4f}'],
    ['확률 → 가중합','dp/dz = p(1 − p) = 0.524979… × 0.475021…', f'{p0*(1-p0):.4f}'],
    ['두 값을 곱함','g_z = (−1/p) × p(1 − p) = p − 1', f'{d:.4f}'],
    ['가중치 w₁','g_w₁ = g_z × x₁ = −0.475021… × 1',f'{d:.4f}'],
    ['가중치 w₂','g_w₂ = g_z × x₂ = −0.475021… × 1',f'{d:.4f}'],
    ['편향 b','g_b = g_z × 1',f'{d:.4f}'],
], [3.2,10.8,3.5])
para('가중합의 기울기는 p − y로 바로 구해도 같습니다. 편향의 미분값은 1이므로 g_b는 g_z와 같아요.', size=9.5)
heading('3 경사하강법으로 한 번 갱신')
table(['대상', '현재 값 − 학습률 × 기울기', '새 값'], [
    ['w₁','0.2 − 0.5 × (−0.475021…)',f'{w1:.4f}'],
    ['w₂','−0.1 − 0.5 × (−0.475021…)',f'{w2:.4f}'],
    ['b','0 − 0.5 × (−0.475021…)',f'{b:.4f}'],
], [3.2,10.8,3.5])
para(f'다시 순전파하면 z = {znew:.4f}, p = {pnew:.4f}, 손실 L = {-math.log(pnew):.4f}입니다. 정답 1의 확률은 올라가고 손실은 줄었습니다.')

# Page 3
d0=p0
newpage('시그모이드와 BCEE 정답이 0인 경우')
para('앞 페이지의 갱신 전 초기값으로 되돌리고 정답만 y = 0으로 바꾼 연습 예제입니다. 입력과 가중치가 같으므로 예측 확률은 여전히 p = 0.524979…입니다.')
heading('1 정답에 따라 남는 손실 항')
equation('L = −[y ln(p) + (1 − y) ln(1 − p)]')
table(['정답', '남는 식', '현재 p에서 손실'], [
    ['y = 1', '−ln(p)',f'{-math.log(p0):.4f}'],
    ['y = 0', '−ln(1 − p)',f'{-math.log(1-p0):.4f}'],
], [3.2,8.8,5.5])
heading('2 역전파에서 기울기 계산')
table(['단계', '미분 결과에 숫자 대입', '결과'], [
    ['손실 → 확률','g_p = 1 ÷ (1 − p)',f'{1/(1-p0):.4f}'],
    ['확률 → 가중합','dp/dz = p(1 − p)',f'{p0*(1-p0):.4f}'],
    ['두 값을 곱함','g_z = [1/(1 − p)] × p(1 − p) = p − 0',f'{d0:.4f}'],
    ['각 가중치와 편향','g_w₁ = g_z × 1, g_w₂ = g_z × 1, g_b = g_z',f'각각 {d0:.4f}'],
], [3.2,10.8,3.5])
heading('3 같은 학습률로 한 번 갱신')
table(['대상', '현재 값 − 0.5 × 기울기', '새 값'], [
    ['w₁','0.2 − 0.5 × 0.524979…',f'{.2-.5*d0:.4f}'],
    ['w₂','−0.1 − 0.5 × 0.524979…',f'{-.1-.5*d0:.4f}'],
    ['b','0 − 0.5 × 0.524979…',f'{-.5*d0:.4f}'],
], [3.2,10.8,3.5])
para('이번에는 가중치와 편향을 낮추어 클래스 1의 확률을 줄이는 방향으로 움직입니다.')
heading('4 여러 데이터라면 손실과 기울기를 각각 평균')
para('두 데이터의 개별 손실이 0.2231과 0.3567이고, 같은 가중치에 대한 개별 기울기가 −0.4와 0.2라고 가정해볼게요. 아래 숫자는 평균 방법을 설명하기 위한 별도 예입니다.', size=9.5)
table(['평균 낼 대상', '계산', '결과'], [
    ['배치 손실','(0.2231 + 0.3567) ÷ 2','0.2899'],
    ['같은 가중치의 기울기','(−0.4 + 0.2) ÷ 2','−0.1000'],
], [5.0,9.0,3.5])
para('평균 손실로 학습하면 각 데이터의 기울기도 평균하여 갱신에 사용합니다. BCEE 손실을 평균 내는 것은 MSE 계산과 별개입니다.', size=9.5)

# Page 4
newpage('소프트맥스와 CCEE 여러 클래스의 경우')
para('5차시 붓꽃 예제의 실제 정답과 확률을 사용합니다. 클래스 순서는 Setosa, Versicolor, Virginica이고, 정답은 Virginica이므로 y = [0, 0, 1]입니다.')
heading('1 정답 클래스의 확률로 손실 계산')
equation('L = −Σ yₖ ln(pₖ)')
table(['앱의 예측', '확률 배열', '정답 확률', 'CCEE 손실'], [
    ['A','[0.12, 0.04, 0.84]','0.84','−ln(0.84) = 0.1744'],
    ['B','[0.25, 0.35, 0.40]','0.40','−ln(0.40) = 0.9163'],
    ['C','[0.70, 0.20, 0.10]','0.10','−ln(0.10) = 2.3026'],
], [2,6.2,2.6,6.7])
para('예측 A를 식에 넣으면 −[0 × ln(0.12) + 0 × ln(0.04) + 1 × ln(0.84)]입니다. 정답 위치의 항만 남습니다.', size=9.5)
heading('2 소프트맥스와 CCEE를 함께 역전파')
para('zₖ는 소프트맥스에 넣기 전의 클래스별 점수입니다. 이 조합의 미분 결과는 g_zₖ = pₖ − yₖ이므로 바로 숫자를 넣으면 됩니다.')
table(['클래스', 'pₖ', 'yₖ', '점수의 기울기 pₖ − yₖ'], [
    ['Setosa','0.12','0','0.12 − 0 = 0.1200'],
    ['Versicolor','0.04','0','0.04 − 0 = 0.0400'],
    ['Virginica','0.84','1','0.84 − 1 = −0.1600'],
], [4.1,2,2,9.4])
para('손실 식에서는 정답 항만 남지만, 소프트맥스가 클래스 점수들을 연결하므로 오답 클래스 점수에도 기울기가 생깁니다.', bold=True, size=9.5)
heading('3 출력층 가중치와 편향의 기울기')
para('출력층에 들어온 은닉 뉴런의 출력 하나를 h = 2라고 둔 연습 예제입니다. 각 클래스 점수에 이 값을 연결하는 가중치 vₖ가 있을 때 g_vₖ = g_zₖ × h입니다.', size=9.5)
table(['클래스', '가중치 기울기', '편향 기울기'], [
    ['Setosa','0.12 × 2 = 0.2400','0.1200'],
    ['Versicolor','0.04 × 2 = 0.0800','0.0400'],
    ['Virginica','−0.16 × 2 = −0.3200','−0.1600'],
], [4.1,8.4,5])
para('8차시 실제 모델은 10개 숫자 클래스에 같은 원리를 적용합니다. 기울기를 구한 뒤에는 Adam으로 갱신합니다. Adam은 기울기의 과거 기록도 사용하므로 위 기울기만으로 실제 갱신량을 정할 수는 없습니다.', size=9.5)

# Page 5
newpage('은닉층이 있을 때 순전파와 역전파')
para('앱에서 사용하는 ReLU 은닉층과 시그모이드 출력층을 연결한 숫자 연습입니다. 구조와 함수는 앱에 맞추고, 기울기가 쉽게 보이도록 값은 새로 정했습니다.')
para('입력 2개 → ReLU 은닉 뉴런 2개 → 시그모이드 출력 1개 → BCEE 손실', bold=True)
para('입력 x₁ = 1, x₂ = 2, 정답 y = 1, 학습률 η = 0.1입니다. 은닉층 가중치는 w, 은닉층에서 출력층으로 연결되는 가중치는 v라고 부를게요.')
heading('1 은닉층부터 출력층까지 순전파')
table(['구할 값', '초기값을 넣은 계산', '결과'], [
    ['은닉 뉴런 1의 입력 a₁','w₁₁x₁ + w₁₂x₂ + b₁ = 1 × 1 + 0 × 2 + 0','1'],
    ['은닉 뉴런 1의 출력 h₁','ReLU(1)','1'],
    ['은닉 뉴런 2의 입력 a₂','w₂₁x₁ + w₂₂x₂ + b₂ = −1 × 1 + 0 × 2 + 0','−1'],
    ['은닉 뉴런 2의 출력 h₂','ReLU(−1)','0'],
    ['출력층 점수 z','v₁h₁ + v₂h₂ + bₒ = 1 × 1 + 1 × 0 − 1','0'],
    ['예측 확률 p','시그모이드(0)','0.5000'],
    ['BCEE 손실 L','−ln(0.5)','0.6931'],
], [4.6,10.4,2.5], font=9)
heading('2 출력층에서 기울기 계산')
table(['구할 기울기', '계산', '결과'], [
    ['출력층 점수 g_z','p − y = 0.5 − 1','−0.5'],
    ['출력층 가중치 g_v₁','g_z × h₁ = −0.5 × 1','−0.5'],
    ['출력층 가중치 g_v₂','g_z × h₂ = −0.5 × 0','0'],
    ['출력층 편향 g_bₒ','g_z','−0.5'],
], [4.6,10.4,2.5], font=9)
heading('3 기울기를 은닉층으로 전달')
para('출력층에서 돌아온 기울기에 연결 가중치 v를 곱하고, ReLU의 기울기를 곱합니다. 여기서는 두 연결 가중치의 갱신 전 값이 모두 1입니다.', size=9.5)
table(['은닉 뉴런', '출력에 대한 기울기 g_h', 'ReLU 기울기', '입력에 대한 기울기 g_a'], [
    ['1','g_z × v₁ = −0.5 × 1','a₁ = 1 > 0이므로 1','−0.5 × 1 = −0.5'],
    ['2','g_z × v₂ = −0.5 × 1','a₂ = −1 < 0이므로 0','−0.5 × 0 = 0'],
], [2,5.2,4.8,5.5], font=9)
para('ReLU 입력이 음수인 두 번째 은닉 뉴런은 이 데이터에서 기울기가 0이 됩니다.', size=9.5)

# Page 6
newpage('은닉층 가중치까지 계산하고 한 번 갱신')
para('앞 페이지에서 구한 은닉층 입력의 기울기는 g_a₁ = −0.5, g_a₂ = 0입니다. 여기에 각 입력값을 곱하면 은닉층 가중치의 기울기가 됩니다.')
heading('1 은닉층 가중치와 편향의 기울기')
table(['대상', '기울기 계산', '결과'], [
    ['w₁₁  입력 1 → 은닉 1','g_a₁ × x₁ = −0.5 × 1','−0.5'],
    ['w₁₂  입력 2 → 은닉 1','g_a₁ × x₂ = −0.5 × 2','−1.0'],
    ['b₁  은닉 1 편향','g_a₁','−0.5'],
    ['w₂₁과 w₂₂ 및 b₂','은닉 2의 g_a₂가 0이므로 모두 0','0'],
], [5.2,9.8,2.5])
heading('2 모든 기울기를 구한 뒤 함께 갱신')
para('새 값 = 현재 값 − 0.1 × 기울기를 적용합니다. 기울기를 전달할 때는 모두 갱신 전 가중치를 사용하고, 계산이 끝난 뒤 값을 바꿉니다.', bold=True, size=9.5)
table(['대상', '현재 값', '기울기', '갱신 계산', '새 값'], [
    ['w₁₁','1','−0.5','1 − 0.1 × (−0.5)','1.05'],
    ['w₁₂','0','−1.0','0 − 0.1 × (−1.0)','0.10'],
    ['b₁','0','−0.5','0 − 0.1 × (−0.5)','0.05'],
    ['w₂₁','−1','0','−1 − 0.1 × 0','−1'],
    ['w₂₂와 b₂','각각 0','각각 0','0 − 0.1 × 0','각각 0'],
    ['v₁','1','−0.5','1 − 0.1 × (−0.5)','1.05'],
    ['v₂','1','0','1 − 0.1 × 0','1'],
    ['bₒ','−1','−0.5','−1 − 0.1 × (−0.5)','−0.95'],
], [2.5,2.3,2.3,7.9,2.5], font=9)
heading('3 다시 순전파해서 손실 확인')
pn=1/(1+math.exp(-.415))
table(['확인할 값', '갱신 후 계산', '결과'], [
    ['은닉 출력 h₁','ReLU(1.05 × 1 + 0.10 × 2 + 0.05)','1.3000'],
    ['은닉 출력 h₂','ReLU(−1 × 1 + 0 × 2 + 0)','0'],
    ['출력층 점수 z','1.05 × 1.30 + 1 × 0 − 0.95','0.4150'],
    ['확률과 손실','p = 시그모이드(0.415), L = −ln(p)',f'p = {pn:.4f}\nL = {-math.log(pn):.4f}'],
], [3.3,10.7,3.5], font=9)
para(f'손실은 0.6931에서 {-math.log(pn):.4f}로 줄었습니다. 출력층뿐 아니라 첫 번째 은닉 뉴런의 가중치도 함께 바뀐 결과입니다.', size=9.5)
para('은닉 뉴런이 여러 출력 뉴런에 연결되면, 돌아오는 기울기를 연결마다 구해 더합니다. 일반식은 g_hⱼ = Σₖ(vₖⱼ × g_zₖ)이며, 그다음 ReLU 기울기를 곱합니다.', size=9)
para('앱 확인 기준 2026년 10월 19일  |  3차시 활성화 함수  4차시 신경망 구조  5차시 손실  6차시 학습  8차시 숫자 분류', size=8)

doc.core_properties.title='앱에서 배우는 역전파 숫자 예제와 공식표'
doc.core_properties.subject='BCEE CCEE ReLU 은닉층 역전파와 가중치 갱신'
doc.core_properties.author=''
doc.core_properties.keywords='BCE, CCE, ReLU, 역전파, 은닉층'
for root in [doc.styles.element, doc.element]:
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.save(OUT)
print(OUT)
print({'bce_before':-math.log(p0),'bce_after':-math.log(pnew),'hidden_after':-math.log(pn),'pages_planned':6})
